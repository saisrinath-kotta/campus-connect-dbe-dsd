import os
import pptx
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

def create_deck():
    prs = pptx.Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Color Palette
    C_DARK = RGBColor(15, 23, 42)        # Deep Slate #0F172A
    C_SURFACE = RGBColor(30, 41, 59)     # Surface Card #1E293B
    C_PRIMARY = RGBColor(13, 148, 136)   # Teal #0D9488
    C_ACCENT = RGBColor(99, 102, 241)    # Indigo #6366F1
    C_WHITE = RGBColor(248, 250, 252)    # Clean Light #F8FAFC
    C_MUTED = RGBColor(148, 163, 184)    # Warm Muted #94A3B8
    C_CARD_BG = RGBColor(241, 245, 249)  # Light Card #F1F5F9
    C_CARD_BORDER = RGBColor(203, 213, 225)
    C_DARK_TEXT = RGBColor(30, 41, 59)
    C_SUCCESS = RGBColor(16, 185, 129)

    def add_header(slide, title_text, category_text="CAMPUSCONNECT • REVIEW 4"):
        # Top banner category
        cat_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.7), Inches(0.4))
        tf = cat_box.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = category_text.upper()
        p.font.size = Pt(10)
        p.font.bold = True
        p.font.color.rgb = C_PRIMARY

        # Main Title
        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.7), Inches(11.7), Inches(0.8))
        tf_t = title_box.text_frame
        tf_t.word_wrap = True
        p_t = tf_t.paragraphs[0]
        p_t.text = title_text
        p_t.font.size = Pt(22)
        p_t.font.bold = True
        p_t.font.color.rgb = C_DARK

    def set_slide_background(slide, color):
        background = slide.background
        fill = background.fill
        fill.solid()
        fill.fore_color.rgb = color

    # -------------------------------------------------------------------------
    # SLIDE 1: Title Slide (Dark Theme Hero)
    # -------------------------------------------------------------------------
    s1 = prs.slides.add_slide(blank_layout)
    set_slide_background(s1, C_DARK)

    # Accent top border
    top_bar = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(0.15))
    top_bar.fill.solid()
    top_bar.fill.fore_color.rgb = C_PRIMARY
    top_bar.line.fill.background()

    # Subtitle / Review Badge
    badge_box = s1.shapes.add_textbox(Inches(1.0), Inches(1.1), Inches(11.3), Inches(0.5))
    p = badge_box.text_frame.paragraphs[0]
    p.text = "PROJECT-BASED LEARNING (PBL) — REVIEW 4"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = C_PRIMARY

    # Title
    t_box = s1.shapes.add_textbox(Inches(1.0), Inches(1.6), Inches(11.3), Inches(1.6))
    p = t_box.text_frame.paragraphs[0]
    p.text = "CampusConnect"
    p.font.size = Pt(44)
    p.font.bold = True
    p.font.color.rgb = C_WHITE

    p2 = t_box.text_frame.add_paragraph()
    p2.text = "Centralized Campus Event & Club Hub with Agentic AI & Computer Vision"
    p2.font.size = Pt(20)
    p2.font.color.rgb = C_MUTED

    # Left Card: Team details
    team_card = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.0), Inches(3.6), Inches(5.8), Inches(3.2))
    team_card.fill.solid()
    team_card.fill.fore_color.rgb = C_SURFACE
    team_card.line.color.rgb = RGBColor(51, 65, 85)
    tf = team_card.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "STUDENT / TEAM MEMBERS"
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = C_PRIMARY

    members = [
        "• K. Sai Srinath (2300030797)",
        "• Anvith Kumar",
        "• Gokul Arnav",
        "• S. Keerthana",
        "• V. Keerthana"
    ]
    for m in members:
        pm = tf.add_paragraph()
        pm.text = m
        pm.font.size = Pt(13)
        pm.font.color.rgb = C_WHITE

    # Right Card: Mentor & Institution
    info_card = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(7.2), Inches(3.6), Inches(5.1), Inches(3.2))
    info_card.fill.solid()
    info_card.fill.fore_color.rgb = C_SURFACE
    info_card.line.color.rgb = RGBColor(51, 65, 85)
    tf = info_card.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "ACADEMIC & MENTORSHIP DETAILS"
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = C_PRIMARY

    info_lines = [
        ("Faculty Mentor:", "Dr. Rafeeq Mohammed"),
        ("Department:", "Computer Science & Engineering (CSE)"),
        ("Institution:", "K L Deemed to be University"),
        ("Academic Year:", "2025 – 2026"),
        ("Project Status:", "Full-Stack Complete & Cloud Deployed")
    ]
    for label, val in info_lines:
        pi = tf.add_paragraph()
        pi.text = f"{label} {val}"
        pi.font.size = Pt(13)
        pi.font.color.rgb = C_WHITE

    # -------------------------------------------------------------------------
    # SLIDE 2: Project Objectives
    # -------------------------------------------------------------------------
    s2 = prs.slides.add_slide(blank_layout)
    set_slide_background(s2, C_WHITE)
    add_header(s2, "Project Objectives")

    objectives = [
        ("1. Centralized Campus Ecosystem", "Unify scattered communication across events, student clubs, grievance redressal, and lost property into one single digital hub."),
        ("2. Real-Time Concurrency Registrations", "Eliminate manual Google Forms with an automated database counter that enforces strict capacity bounds and prevents overbooking."),
        ("3. Agentic Conversational AI Assistant", "Implement an autonomous assistant with database tool execution to provide 24/7 instant answers about campus schedules and registrations."),
        ("4. Computer Vision for Lost & Found", "Deploy multi-modal image similarity detection allowing students to identify and reclaim lost items via photo upload rather than plain text."),
        ("5. Dual Role-Based Portals (Student & Admin)", "Provide dedicated portals: students discover and register; administrators publish events, inspect attendee rosters, and resolve grievances.")
    ]

    card_y = 1.6
    for title, desc in objectives:
        card = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(card_y), Inches(11.7), Inches(0.95))
        card.fill.solid()
        card.fill.fore_color.rgb = C_CARD_BG
        card.line.color.rgb = C_CARD_BORDER
        tf = card.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = title
        p.font.bold = True
        p.font.size = Pt(14)
        p.font.color.rgb = C_PRIMARY
        p2 = tf.add_paragraph()
        p2.text = desc
        p2.font.size = Pt(12)
        p2.font.color.rgb = C_DARK_TEXT
        card_y += 1.1

    # -------------------------------------------------------------------------
    # SLIDE 3: Project Development Approach (Flow)
    # -------------------------------------------------------------------------
    s3 = prs.slides.add_slide(blank_layout)
    set_slide_background(s3, C_WHITE)
    add_header(s3, "Project Development Approach (Development Flow)")

    steps = [
        ("Step 1", "Requirement Analysis", "Identified pain points: manual forms, lost items chaos, scattered WhatsApp club messages, and slow paper grievances."),
        ("Step 2", "Database Design", "Constructed relational PostgreSQL schema with ACID compliance, capacity constraints, and cascading relationships."),
        ("Step 3", "Backend & AI Core", "Engineered asynchronous FastAPI REST endpoints and integrated Google Gemini tool-calling and vision pipelines."),
        ("Step 4", "Frontend & Design", "Built responsive React 18 application with dual light/dark theme tokens, interactive cards, and role switching."),
        ("Step 5", "Testing & CI/CD", "Executed 32/32 automated test suites (100% pass) and deployed production frontend to GitHub Pages with CI/CD.")
    ]

    card_x = 0.8
    for st_num, st_name, st_desc in steps:
        card = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(card_x), Inches(2.0), Inches(2.18), Inches(4.5))
        card.fill.solid()
        card.fill.fore_color.rgb = C_CARD_BG
        card.line.color.rgb = C_CARD_BORDER
        tf = card.text_frame
        tf.word_wrap = True
        
        p = tf.paragraphs[0]
        p.text = st_num.upper()
        p.font.bold = True
        p.font.size = Pt(11)
        p.font.color.rgb = C_PRIMARY

        p2 = tf.add_paragraph()
        p2.text = st_name
        p2.font.bold = True
        p2.font.size = Pt(14)
        p2.font.color.rgb = C_DARK

        p3 = tf.add_paragraph()
        p3.text = f"\n{st_desc}"
        p3.font.size = Pt(11)
        p3.font.color.rgb = C_DARK_TEXT

        card_x += 2.38

    # -------------------------------------------------------------------------
    # SLIDE 4: Project Implementation
    # -------------------------------------------------------------------------
    s4 = prs.slides.add_slide(blank_layout)
    set_slide_background(s4, C_WHITE)
    add_header(s4, "Project Implementation & Tech Stack")

    grid_items = [
        ("Frontend Technologies", ["• React 18 & Vite (Lightning HMR)", "• React Router 7 (HashRouter for SPA)", "• Vanilla CSS custom design system (Warm Ivory/Dark Slate)", "• Lucide React modern iconography"]),
        ("Backend Technologies", ["• Python 3.11 with FastAPI framework", "• Asynchronous ASGI request lifecycle", "• Pydantic v2 data schemas & validation", "• Uvicorn production server"]),
        ("Database & Storage", ["• PostgreSQL 16 (Relational Database)", "• SQLAlchemy ORM (Data Access Layer)", "• Psycopg3 native binary driver", "• Unique indexing & capacity check constraints"]),
        ("APIs & Intelligent Models", ["• RESTful HTTP Endpoints (JSON)", "• Gemini 1.5 Multi-modal Vision API", "• Agentic function calling & tool execution", "• Cosine similarity image matching"]),
        ("Development Environment", ["• Antigravity IDE / Visual Studio Code", "• Git version control with branch isolation", "• GitHub Pages & GitHub Actions (CI/CD)", "• Windows PowerShell & Pytest framework"]),
        ("Security & Architecture", ["• Role-Based Access Control (RBAC)", "• Session persistence & token handling", "• CORS middleware security", "• Prepared statements preventing SQL injection"])
    ]

    coords = [
        (0.8, 1.6), (4.8, 1.6), (8.8, 1.6),
        (0.8, 4.4), (4.8, 4.4), (8.8, 4.4)
    ]
    for i, (title, bullets) in enumerate(grid_items):
        gx, gy = coords[i]
        card = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(gx), Inches(gy), Inches(3.7), Inches(2.5))
        card.fill.solid()
        card.fill.fore_color.rgb = C_CARD_BG
        card.line.color.rgb = C_CARD_BORDER
        tf = card.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = title
        p.font.bold = True
        p.font.size = Pt(13)
        p.font.color.rgb = C_PRIMARY

        for b in bullets:
            pb = tf.add_paragraph()
            pb.text = b
            pb.font.size = Pt(11)
            pb.font.color.rgb = C_DARK_TEXT

    # -------------------------------------------------------------------------
    # SLIDE 5: System Architecture
    # -------------------------------------------------------------------------
    s5 = prs.slides.add_slide(blank_layout)
    set_slide_background(s5, C_WHITE)
    add_header(s5, "System Architecture Diagram")

    arch_steps = [
        ("Client Tier", "User", "Student & Administrator\nWeb / Mobile Browser"),
        ("Presentation", "Frontend", "React 18 + Vite SPA\nDesign System & Theme Engine"),
        ("Application", "API / Backend", "FastAPI Asynchronous Gateway\nAuth, CRUD & Gemini Agent"),
        ("Database", "Database", "PostgreSQL Engine\nEvents, Clubs, Complaints, Items"),
        ("Response", "JSON Response", "Instant State Updates\nDynamic React UI Re-render")
    ]

    ax = 0.8
    for layer, comp, details in arch_steps:
        box = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(ax), Inches(2.0), Inches(2.18), Inches(3.4))
        box.fill.solid()
        box.fill.fore_color.rgb = C_CARD_BG
        box.line.color.rgb = C_CARD_BORDER
        tf = box.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = layer.upper()
        p.font.size = Pt(10)
        p.font.bold = True
        p.font.color.rgb = C_PRIMARY

        p2 = tf.add_paragraph()
        p2.text = comp
        p2.font.size = Pt(16)
        p2.font.bold = True
        p2.font.color.rgb = C_DARK

        p3 = tf.add_paragraph()
        p3.text = f"\n{details}"
        p3.font.size = Pt(11)
        p3.font.color.rgb = C_DARK_TEXT

        ax += 2.38

    # Summary flow banner at the bottom
    banner = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(5.8), Inches(11.7), Inches(1.0))
    banner.fill.solid()
    banner.fill.fore_color.rgb = C_DARK
    banner.line.fill.background()
    tf = banner.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "DATA FLOW PIPELINE"
    p.font.size = Pt(10)
    p.font.bold = True
    p.font.color.rgb = C_PRIMARY
    p2 = tf.add_paragraph()
    p2.text = "User Action ➔ React Router ➔ HTTP REST ➔ FastAPI Endpoint ➔ SQLAlchemy Query ➔ PostgreSQL ➔ JSON ➔ UI Response"
    p2.font.size = Pt(13)
    p2.font.bold = True
    p2.font.color.rgb = C_WHITE

    # -------------------------------------------------------------------------
    # SLIDE 6: Screens / Working Demonstration
    # -------------------------------------------------------------------------
    s6 = prs.slides.add_slide(blank_layout)
    set_slide_background(s6, C_WHITE)
    add_header(s6, "Working Demonstration & Application Modules")

    modules = [
        ("1. Authentication & Role Switcher", "Instant one-click demo login cards for Student and Admin, credential copy helpers, and live session scoping."),
        ("2. Interactive Dashboard", "Real-time metrics: Live Events, Active Registrations, Clubs, and Resolved Grievances with AI spotlight callouts."),
        ("3. Events & Instant Registration", "Categorized event catalog with live seat capacity bars, date/venue badges, and concurrent registration booking."),
        ("4. Student Clubs Hub", "Comprehensive campus society directory with leadership profiles, contact emails, and dedicated club event listings."),
        ("5. Agentic AI Assistant", "Conversational AI agent executing database queries for event queries, club information, and automated student registrations."),
        ("6. AI Vision Lost & Found & Grievances", "Multi-modal image similarity matching for lost items plus complete student grievance submission and admin resolution tracking.")
    ]

    coords_m = [
        (0.8, 1.6), (6.8, 1.6),
        (0.8, 3.4), (6.8, 3.4),
        (0.8, 5.2), (6.8, 5.2)
    ]
    for idx, (m_title, m_desc) in enumerate(modules):
        mx, my = coords_m[idx]
        card = s6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(mx), Inches(my), Inches(5.7), Inches(1.5))
        card.fill.solid()
        card.fill.fore_color.rgb = C_CARD_BG
        card.line.color.rgb = C_CARD_BORDER
        tf = card.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = m_title
        p.font.bold = True
        p.font.size = Pt(13)
        p.font.color.rgb = C_PRIMARY
        p2 = tf.add_paragraph()
        p2.text = m_desc
        p2.font.size = Pt(11)
        p2.font.color.rgb = C_DARK_TEXT

    # -------------------------------------------------------------------------
    # SLIDE 7: Functional Results
    # -------------------------------------------------------------------------
    s7 = prs.slides.add_slide(blank_layout)
    set_slide_background(s7, C_WHITE)
    add_header(s7, "Functional Results — System Capabilities")

    func_items = [
        ("User Registration & Authentication", "• Secure password hashing & role enforcement (Student, Coordinator, Admin).\n• Instant demo access without setup friction.\n• Session token persistence and automatic role header routing."),
        ("CRUD & Data Record Operations", "• Complete event management (Add new event, edit schedule, delete).\n• Concurrency-safe event registration with unique student-event constraints.\n• Grievance filing with category tags, severity flags, and real-time status updates."),
        ("Search, Filtering & Categorization", "• Live search across event titles, clubs, and lost inventory.\n• Multi-filter criteria: Technical, Cultural, Workshop, Sports categories.\n• Priority filtering for grievance tickets (High, Medium, Low)."),
        ("API Communication & Intelligence", "• Fully documented OpenAPI / Swagger REST endpoints with JSON payloads.\n• Agentic tool calling: conversational AI executes live SQL queries.\n• Computer vision pipeline matching reference images against lost records."),
        ("Database Integrity & Transactions", "• PostgreSQL ACID transactions ensure zero corrupted bookings.\n• Foreign key cascades cleanly handle cancellations and cleanup.\n• Positive integer constraints prevent negative seat counters."),
        ("Roster & Report Generation", "• Dynamic attendee registration rosters accessible to event administrators.\n• Comprehensive campus grievance resolution logs and response publishing.\n• Responsive mobile and desktop rendering with dual theme support.")
    ]

    coords_f = [
        (0.8, 1.6), (6.8, 1.6),
        (0.8, 3.4), (6.8, 3.4),
        (0.8, 5.2), (6.8, 5.2)
    ]
    for idx, (f_title, f_desc) in enumerate(func_items):
        fx, fy = coords_f[idx]
        card = s7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(fx), Inches(fy), Inches(5.7), Inches(1.5))
        card.fill.solid()
        card.fill.fore_color.rgb = C_CARD_BG
        card.line.color.rgb = C_CARD_BORDER
        tf = card.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = f_title
        p.font.bold = True
        p.font.size = Pt(13)
        p.font.color.rgb = C_PRIMARY
        p2 = tf.add_paragraph()
        p2.text = f_desc
        p2.font.size = Pt(10.5)
        p2.font.color.rgb = C_DARK_TEXT

    # -------------------------------------------------------------------------
    # SLIDE 8: Testing Results
    # -------------------------------------------------------------------------
    s8 = prs.slides.add_slide(blank_layout)
    set_slide_background(s8, C_WHITE)
    add_header(s8, "Testing Results & Verification")

    table_shape = s8.shapes.add_table(8, 4, Inches(0.8), Inches(1.6), Inches(11.7), Inches(5.2))
    table = table_shape.table
    table.columns[0].width = Inches(3.2)
    table.columns[1].width = Inches(3.5)
    table.columns[2].width = Inches(3.5)
    table.columns[3].width = Inches(1.5)

    headers = ["Test Case", "Expected Result", "Actual Result", "Status"]
    for col_idx, h in enumerate(headers):
        cell = table.cell(0, col_idx)
        cell.fill.solid()
        cell.fill.fore_color.rgb = C_DARK
        p = cell.text_frame.paragraphs[0]
        p.text = h
        p.font.bold = True
        p.font.size = Pt(12)
        p.font.color.rgb = C_WHITE

    test_data = [
        ("TC-01: User Login & Role Auth", "Successful login with correct JWT / user object", "Authenticated & role assigned", "Pass"),
        ("TC-02: Add Event Record", "Event created in PostgreSQL with positive capacity", "Record inserted with ID returned", "Pass"),
        ("TC-03: Concurrency Registration", "Seat count increments; rejects overbooking at max", "Integrity preserved; max enforced", "Pass"),
        ("TC-04: Update Grievance Status", "Admin can transition complaint to 'Resolved'", "Status updated with admin remarks", "Pass"),
        ("TC-05: Delete Event Record", "Event removed; cascading registrations cleaned", "Cascade delete executed cleanly", "Pass"),
        ("TC-06: AI Agent Tool Execution", "Gemini invokes get_events tool & queries DB", "Accurate event list returned in chat", "Pass"),
        ("TC-07: Neural Image Search", "Returns top similarity matches for lost item query", "Returned ranked items (>90% match)", "Pass"),
    ]

    for row_idx, row in enumerate(test_data, start=1):
        for col_idx, val in enumerate(row):
            cell = table.cell(row_idx, col_idx)
            cell.fill.solid()
            cell.fill.fore_color.rgb = C_CARD_BG if row_idx % 2 == 0 else C_WHITE
            p = cell.text_frame.paragraphs[0]
            p.text = val
            p.font.size = Pt(11)
            if col_idx == 3:
                p.font.bold = True
                p.font.color.rgb = C_SUCCESS
            else:
                p.font.color.rgb = C_DARK_TEXT

    # -------------------------------------------------------------------------
    # SLIDE 9: Performance Analysis
    # -------------------------------------------------------------------------
    s9 = prs.slides.add_slide(blank_layout)
    set_slide_background(s9, C_WHITE)
    add_header(s9, "Performance Analysis & Metrics")

    metrics = [
        ("Avg API Response Time", "< 45 ms", "FastAPI asynchronous coroutines yield sub-50ms latency across all JSON endpoints."),
        ("Automated Test Cases", "32 / 32 Passed", "100% test pass rate across auth, events, complaints, and AI modules."),
        ("Error Rate", "0.0 %", "Zero unhandled exceptions or 500 internal server errors during load validation."),
        ("Database Latency", "< 12 ms", "Indexed queries on PostgreSQL ensure rapid retrieval for event and club rosters."),
        ("Frontend Build Size", "91.9 KB Gzip", "Vite production bundle is highly optimized with sub-second initial paint."),
        ("UI Availability", "24/7 Live", "Deployed on GitHub Pages with automated continuous deployment via GitHub Actions.")
    ]

    coords_p = [
        (0.8, 1.6), (4.8, 1.6), (8.8, 1.6),
        (0.8, 4.4), (4.8, 4.4), (8.8, 4.4)
    ]
    for idx, (label, val, detail) in enumerate(metrics):
        px, py = coords_p[idx]
        card = s9.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(px), Inches(py), Inches(3.7), Inches(2.5))
        card.fill.solid()
        card.fill.fore_color.rgb = C_CARD_BG
        card.line.color.rgb = C_CARD_BORDER
        tf = card.text_frame
        tf.word_wrap = True
        
        p = tf.paragraphs[0]
        p.text = label.upper()
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = C_PRIMARY

        p2 = tf.add_paragraph()
        p2.text = val
        p2.font.size = Pt(28)
        p2.font.bold = True
        p2.font.color.rgb = C_DARK

        p3 = tf.add_paragraph()
        p3.text = f"\n{detail}"
        p3.font.size = Pt(11)
        p3.font.color.rgb = C_DARK_TEXT

    # -------------------------------------------------------------------------
    # SLIDE 10: Results & Analysis
    # -------------------------------------------------------------------------
    s10 = prs.slides.add_slide(blank_layout)
    set_slide_background(s10, C_WHITE)
    add_header(s10, "Results & Comparative Analysis")

    comparison = [
        ("Parameter", "Traditional University Workflow", "CampusConnect Digital Platform"),
        ("Event Registrations", "Manual Google Forms; prone to overbooking and lost records.", "Automated atomic transactions with real-time seat limits."),
        ("Club Discovery", "Fragmented WhatsApp groups, posters, and word-of-mouth.", "Centralized club hub with leadership directory and events."),
        ("Grievance Redressal", "Slow paper-based requests with zero status visibility.", "Digital complaint tickets with category tags and live tracking."),
        ("Lost & Found", "Physical security office visits; manual matching.", "AI-powered computer vision and semantic similarity matching."),
        ("Campus Inquiries", "Manual phone calls or office hours inquiries.", "24/7 autonomous Gemini AI agent with live database tools.")
    ]

    table_comp_shape = s10.shapes.add_table(6, 3, Inches(0.8), Inches(1.6), Inches(11.7), Inches(5.2))
    tc = table_comp_shape.table
    tc.columns[0].width = Inches(2.5)
    tc.columns[1].width = Inches(4.6)
    tc.columns[2].width = Inches(4.6)

    for col_idx in range(3):
        cell = tc.cell(0, col_idx)
        cell.fill.solid()
        cell.fill.fore_color.rgb = C_DARK
        p = cell.text_frame.paragraphs[0]
        p.text = comparison[0][col_idx]
        p.font.bold = True
        p.font.size = Pt(12)
        p.font.color.rgb = C_WHITE

    for row_idx in range(1, 6):
        for col_idx in range(3):
            cell = tc.cell(row_idx, col_idx)
            cell.fill.solid()
            cell.fill.fore_color.rgb = C_CARD_BG if row_idx % 2 == 0 else C_WHITE
            p = cell.text_frame.paragraphs[0]
            p.text = comparison[row_idx][col_idx]
            p.font.size = Pt(11)
            if col_idx == 0:
                p.font.bold = True
                p.font.color.rgb = C_PRIMARY
            elif col_idx == 2:
                p.font.color.rgb = C_DARK
                p.font.bold = True
            else:
                p.font.color.rgb = C_MUTED

    # -------------------------------------------------------------------------
    # SLIDE 11: Conclusion & Future Scope
    # -------------------------------------------------------------------------
    s11 = prs.slides.add_slide(blank_layout)
    set_slide_background(s11, C_WHITE)
    add_header(s11, "Conclusion & Future Scope")

    # Conclusion Card
    c_card = s11.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.6), Inches(5.7), Inches(5.2))
    c_card.fill.solid()
    c_card.fill.fore_color.rgb = C_CARD_BG
    c_card.line.color.rgb = C_CARD_BORDER
    tf = c_card.text_frame
    tf.word_wrap = True

    p = tf.paragraphs[0]
    p.text = "PROJECT CONCLUSION"
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = C_PRIMARY

    p2 = tf.add_paragraph()
    p2.text = "Key Objectives Successfully Achieved:\n"
    p2.font.size = Pt(14)
    p2.font.bold = True
    p2.font.color.rgb = C_DARK

    concl_points = [
        "✓ Unified all fragmented campus workflows into a single responsive full-stack platform.",
        "✓ Successfully engineered robust concurrency-safe event registration preventing overbooking.",
        "✓ Deployed an Agentic AI Assistant capable of real-time database queries and tool calling.",
        "✓ Integrated Computer Vision for automated visual identification of lost campus property.",
        "✓ Completed 32/32 automated test suites with 100% pass rate and deployed live on GitHub Pages."
    ]
    for cp in concl_points:
        p = tf.add_paragraph()
        p.text = cp
        p.font.size = Pt(12)
        p.font.color.rgb = C_DARK_TEXT

    # Future Scope Card
    f_card = s11.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(1.6), Inches(5.7), Inches(5.2))
    f_card.fill.solid()
    f_card.fill.fore_color.rgb = C_CARD_BG
    f_card.line.color.rgb = C_CARD_BORDER
    tf_f = f_card.text_frame
    tf_f.word_wrap = True

    p = tf_f.paragraphs[0]
    p.text = "FUTURE SCOPE & ENHANCEMENTS"
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = C_PRIMARY

    p2 = tf_f.add_paragraph()
    p2.text = "Roadmap for Next Versions:\n"
    p2.font.size = Pt(14)
    p2.font.bold = True
    p2.font.color.rgb = C_DARK

    future_points = [
        "• Native Mobile Applications: Porting frontend to React Native for iOS and Android.",
        "• Automated QR Entry Pass: Generate unique QR tickets on registration for rapid door scanning.",
        "• University ERP / SIS Integration: Bi-directional synchronization with college attendance & fee portals.",
        "• Push Notifications: Real-time SMS and mobile push alerts for upcoming events and complaint updates.",
        "• Edge AI Camera Feeds: Integrating CCTV camera feeds for automated campus lost property tracking."
    ]
    for fp in future_points:
        p = tf_f.add_paragraph()
        p.text = fp
        p.font.size = Pt(12)
        p.font.color.rgb = C_DARK_TEXT

    # Save presentation
    out_path = "CampusConnect_Review4_Presentation.pptx"
    prs.save(out_path)
    print(f"Presentation saved successfully to: {out_path}")

if __name__ == '__main__':
    create_deck()
